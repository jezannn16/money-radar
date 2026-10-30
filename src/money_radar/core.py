"""Money domain, CSV/XLSX preview-import, reports and CBR exchange rates."""
from __future__ import annotations

import csv
import hashlib
import io
import os
import re
import sqlite3
import ssl
import urllib.error
import urllib.request
import xml.etree.ElementTree as ET
import zipfile
from collections.abc import Callable
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_EVEN, Decimal, InvalidOperation
from pathlib import Path

import certifi
from openpyxl.utils.exceptions import InvalidFileException


def money(value: str | Decimal, currency: str) -> Decimal:
    if len(currency) != 3 or not currency.isalpha():
        raise ValueError("Currency must be a three-letter ISO code")
    try:
        amount = Decimal(str(value))
    except InvalidOperation as exc:
        raise ValueError("Amount is not a valid decimal") from exc
    if not amount.is_finite() or amount <= 0:
        raise ValueError("Expense amount must be a finite positive number")
    return amount.quantize(Decimal("0.01"), rounding=ROUND_HALF_EVEN)


@dataclass(frozen=True)
class TransactionInput:
    booked_on: date
    amount: Decimal
    currency: str
    merchant: str
    category: str = "Uncategorized"
    description: str = ""
    source: str = "manual"
    external_id: str | None = None


class Ledger:
    def __init__(self, path: str | Path):
        self.path = str(path)
        with self._db() as db:
            db.executescript("""
            CREATE TABLE IF NOT EXISTS categories(id INTEGER PRIMARY KEY,user_id INTEGER NOT NULL,name TEXT NOT NULL,UNIQUE(user_id,name));
            CREATE TABLE IF NOT EXISTS transactions(
              id INTEGER PRIMARY KEY,user_id INTEGER NOT NULL,booked_on TEXT NOT NULL,amount TEXT NOT NULL,
              currency TEXT NOT NULL,merchant TEXT NOT NULL,category TEXT NOT NULL,description TEXT NOT NULL DEFAULT '',
              source TEXT NOT NULL,external_id TEXT,fingerprint TEXT NOT NULL,import_batch_id INTEGER,
              created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP);
            CREATE UNIQUE INDEX IF NOT EXISTS uq_transaction_external ON transactions(user_id,source,external_id) WHERE external_id IS NOT NULL;
            CREATE UNIQUE INDEX IF NOT EXISTS uq_transaction_fingerprint ON transactions(user_id,fingerprint);
            CREATE TABLE IF NOT EXISTS monthly_budgets(user_id INTEGER NOT NULL,month TEXT NOT NULL,category TEXT NOT NULL,amount TEXT NOT NULL,currency TEXT NOT NULL,PRIMARY KEY(user_id,month,category,currency));
            CREATE TABLE IF NOT EXISTS recurring_rules(id INTEGER PRIMARY KEY,user_id INTEGER NOT NULL,merchant TEXT NOT NULL,category TEXT NOT NULL,amount TEXT,currency TEXT,enabled INTEGER NOT NULL DEFAULT 1,day_of_month INTEGER NOT NULL DEFAULT 1);
            CREATE TABLE IF NOT EXISTS recurring_executions(rule_id INTEGER NOT NULL,period TEXT NOT NULL,state TEXT NOT NULL,transaction_id INTEGER,created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,PRIMARY KEY(rule_id,period));
            CREATE TABLE IF NOT EXISTS import_batches(id INTEGER PRIMARY KEY,user_id INTEGER NOT NULL,file_hash TEXT NOT NULL,created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,UNIQUE(user_id,file_hash));
            CREATE TABLE IF NOT EXISTS exchange_rates(currency TEXT NOT NULL,requested_date TEXT NOT NULL,effective_date TEXT NOT NULL,rate TEXT NOT NULL,fetched_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,PRIMARY KEY(currency,requested_date));
            CREATE TABLE IF NOT EXISTS schema_migrations(version INTEGER PRIMARY KEY,applied_at TEXT NOT NULL);
            INSERT OR IGNORE INTO schema_migrations(version,applied_at) VALUES(1,CURRENT_TIMESTAMP);
            """)
            columns = {row[1] for row in db.execute("PRAGMA table_info(recurring_rules)")}
            if "day_of_month" not in columns:
                db.execute("ALTER TABLE recurring_rules ADD COLUMN day_of_month INTEGER NOT NULL DEFAULT 1")
            db.execute("INSERT OR IGNORE INTO schema_migrations(version,applied_at) VALUES(2,CURRENT_TIMESTAMP)")
            db.execute("INSERT OR IGNORE INTO schema_migrations(version,applied_at) VALUES(3,CURRENT_TIMESTAMP)")

    def _db(self):
        if self.path != ":memory:":
            fd=os.open(self.path,os.O_CREAT|os.O_RDWR,0o600);os.close(fd)
        db=sqlite3.connect(self.path,timeout=10)
        if self.path != ":memory:":
            for candidate in (self.path,self.path+"-wal",self.path+"-shm",self.path+"-journal"):
                try:os.chmod(candidate,0o600)
                except FileNotFoundError:pass
        db.row_factory=sqlite3.Row; return db

    @staticmethod
    def fingerprint(row: TransactionInput) -> str:
        # Deliberately includes date, currency and merchant; equal amounts alone are not duplicates.
        canonical="|".join((row.booked_on.isoformat(),str(row.amount),row.currency.upper(),row.merchant.strip().casefold()))
        return hashlib.sha256(canonical.encode()).hexdigest()

    def add(self,user_id:int,row:TransactionInput,batch_id:int|None=None)->bool:
        if not row.merchant.strip(): raise ValueError("Merchant cannot be empty")
        amount=money(row.amount,row.currency)
        try:
            with self._db() as db:
                db.execute("INSERT INTO transactions(user_id,booked_on,amount,currency,merchant,category,description,source,external_id,fingerprint,import_batch_id) VALUES(?,?,?,?,?,?,?,?,?,?,?)",
                    (user_id,row.booked_on.isoformat(),str(amount),row.currency.upper(),row.merchant.strip(),row.category.strip() or "Uncategorized",row.description,row.source,row.external_id,self.fingerprint(row),batch_id))
            return True
        except sqlite3.IntegrityError:
            return False

    def add_category(self,user_id:int,name:str)->None:
        name=name.strip()
        if not name: raise ValueError("Category name cannot be empty")
        with self._db() as db: db.execute("INSERT OR IGNORE INTO categories(user_id,name) VALUES(?,?)",(user_id,name))

    def categories(self,user_id:int)->list[str]:
        with self._db() as db:
            return [row[0] for row in db.execute("SELECT name FROM categories WHERE user_id=? UNION SELECT category FROM transactions WHERE user_id=? ORDER BY 1 COLLATE NOCASE",(user_id,user_id))]

    def filter_transactions(
        self,user_id:int,*,date_from:str|None=None,date_to:str|None=None,
        category:str|None=None,currency:str|None=None,merchant:str|None=None,
        limit:int=100,
    )->list[sqlite3.Row]:
        clauses=["user_id=?"]; values:list[object]=[user_id]
        for field,value,operator in (("booked_on",date_from,">="),("booked_on",date_to,"<=")):
            if value:
                try: date.fromisoformat(value)
                except ValueError as exc: raise ValueError("Dates must use YYYY-MM-DD") from exc
                clauses.append(f"{field}{operator}?"); values.append(value)
        if date_from and date_to and date_from>date_to:
            raise ValueError("date_from must be on or before date_to")
        for column,value in (("category",category),("currency",currency)):
            if value:
                clauses.append(f"{column}=?"); values.append(value.strip() if column=="category" else value.strip().upper())
        if merchant:
            clauses.append("merchant LIKE ? COLLATE NOCASE"); values.append(f"%{merchant.strip()}%")
        if not 1<=limit<=500:
            raise ValueError("Limit must be between 1 and 500")
        query=("SELECT id,booked_on,amount,currency,merchant,category,description,source,external_id "
               "FROM transactions WHERE "+" AND ".join(clauses)+" ORDER BY booked_on DESC,id DESC LIMIT ?")
        values.append(limit)
        with self._db() as db:return list(db.execute(query,values))

    def cached_exchange_rate(self,currency:str,requested_date:date)->tuple[Decimal,date]|None:
        with self._db() as db:
            row=db.execute("SELECT rate,effective_date FROM exchange_rates WHERE currency=? AND requested_date=?",(currency,requested_date.isoformat())).fetchone()
        return (Decimal(row["rate"]),date.fromisoformat(row["effective_date"])) if row else None

    def save_exchange_rate(self,currency:str,requested_date:date,effective_date:date,rate:Decimal)->None:
        if not rate.is_finite() or rate<=0:raise ValueError("Exchange rate must be positive and finite")
        with self._db() as db:
            db.execute("INSERT INTO exchange_rates(currency,requested_date,effective_date,rate) VALUES(?,?,?,?) ON CONFLICT(currency,requested_date) DO UPDATE SET effective_date=excluded.effective_date,rate=excluded.rate,fetched_at=CURRENT_TIMESTAMP",(currency,requested_date.isoformat(),effective_date.isoformat(),str(rate)))

    def add_recurring_rule(
        self, user_id: int, merchant: str, category: str, amount: Decimal,
        currency: str, day_of_month: int = 1,
    ) -> int:
        merchant = merchant.strip()
        category = category.strip() or "Uncategorized"
        if not merchant:
            raise ValueError("Merchant cannot be empty")
        if not 1 <= day_of_month <= 31:
            raise ValueError("Day of month must be between 1 and 31")
        normalized_amount = money(amount, currency)
        with self._db() as db:
            cur = db.execute(
                "INSERT INTO recurring_rules(user_id,merchant,category,amount,currency,day_of_month) VALUES(?,?,?,?,?,?)",
                (user_id, merchant, category, str(normalized_amount), currency.upper(), day_of_month),
            )
            return int(cur.lastrowid)

    def recurring_rules(self, user_id: int) -> list[sqlite3.Row]:
        with self._db() as db:
            return list(db.execute(
                "SELECT id,merchant,category,amount,currency,day_of_month,enabled FROM recurring_rules WHERE user_id=? ORDER BY id",
                (user_id,),
            ))

    def disable_recurring_rule(self, user_id: int, rule_id: int) -> bool:
        with self._db() as db:
            cur = db.execute(
                "UPDATE recurring_rules SET enabled=0 WHERE id=? AND user_id=? AND enabled=1",
                (rule_id, user_id),
            )
            return bool(cur.rowcount)

    def run_recurring(self, user_id: int, period: str) -> dict[str, int]:
        """Apply enabled monthly rules once; rule/period uniqueness makes reruns idempotent."""
        if not re.fullmatch(r"\d{4}-\d{2}", period):
            raise ValueError("Period must use YYYY-MM")
        try:
            year, month_num = (int(part) for part in period.split("-"))
            first_day = date(year, month_num, 1)
        except ValueError as exc:
            raise ValueError("Period must use YYYY-MM") from exc
        next_month = date(year + (month_num == 12), month_num % 12 + 1, 1)
        last_day = (next_month - timedelta(days=1)).day
        created = duplicates = already_processed = 0
        with self._db() as db:
            rules = list(db.execute(
                "SELECT id,merchant,category,amount,currency,day_of_month FROM recurring_rules WHERE user_id=? AND enabled=1 AND amount IS NOT NULL AND currency IS NOT NULL ORDER BY id",
                (user_id,),
            ))
            for rule in rules:
                cur = db.execute(
                    "INSERT OR IGNORE INTO recurring_executions(rule_id,period,state) VALUES(?,?,'running')",
                    (rule["id"], period),
                )
                if not cur.rowcount:
                    already_processed += 1
                    continue
                booked_on = first_day.replace(day=min(rule["day_of_month"], last_day))
                row = TransactionInput(
                    booked_on, Decimal(rule["amount"]), rule["currency"],
                    rule["merchant"], rule["category"], source="recurring",
                    external_id=f"recurring:{rule['id']}:{period}",
                )
                try:
                    tx = db.execute(
                        "INSERT INTO transactions(user_id,booked_on,amount,currency,merchant,category,description,source,external_id,fingerprint) VALUES(?,?,?,?,?,?,?,?,?,?)",
                        (user_id, booked_on.isoformat(), str(money(row.amount, row.currency)), row.currency.upper(), row.merchant, row.category, "Monthly recurring rule", row.source, row.external_id, self.fingerprint(row)),
                    )
                except sqlite3.IntegrityError:
                    db.execute("UPDATE recurring_executions SET state='duplicate' WHERE rule_id=? AND period=?", (rule["id"], period))
                    duplicates += 1
                else:
                    db.execute("UPDATE recurring_executions SET state='created',transaction_id=? WHERE rule_id=? AND period=?", (tx.lastrowid, rule["id"], period))
                    created += 1
        return {"created": created, "duplicates": duplicates, "already_processed": already_processed}

    def set_budget(self,user_id:int,month:str,category:str,amount:Decimal,currency:str)->None:
        try: datetime.strptime(month,"%Y-%m")
        except ValueError as exc: raise ValueError("Month must use YYYY-MM") from exc
        amt=money(amount,currency)
        with self._db() as db:
            db.execute("INSERT INTO monthly_budgets VALUES(?,?,?,?,?) ON CONFLICT(user_id,month,category,currency) DO UPDATE SET amount=excluded.amount",(user_id,month,category,str(amt),currency.upper()))

    def month_report(self,user_id:int,month:str)->dict:
        with self._db() as db:
            rows=list(db.execute("SELECT category,currency,amount FROM transactions WHERE user_id=? AND substr(booked_on,1,7)=? ORDER BY category",(user_id,month)))
            budgets=list(db.execute("SELECT category,currency,amount FROM monthly_budgets WHERE user_id=? AND month=?",(user_id,month)))
        totals={}
        for row in rows:
            key=(row["category"],row["currency"])
            spent,count=totals.get(key,(Decimal("0.00"),0)); totals[key]=(spent+Decimal(row["amount"]),count+1)
        report=[]
        for (category,currency),(spent,count) in sorted(totals.items()):
            budget=next((Decimal(x["amount"]) for x in budgets if x["category"]==category and x["currency"]==currency),None)
            report.append({"category":category,"currency":currency,"spent":spent,"count":count,"budget":budget,"remaining":budget-spent if budget is not None else None})
        return {"month":month,"items":report,"budget_categories_without_spending":[{"category":b["category"],"currency":b["currency"],"budget":Decimal(b["amount"])} for b in budgets if not any(i["category"]==b["category"] and i["currency"]==b["currency"] for i in report)]}

    def unusual_expenses(
        self, user_id: int, month: str, lookback_months: int = 6,
        factor: Decimal = Decimal(3), minimum_samples: int = 5,
    ) -> list[dict]:
        """Flag transparent candidates above a historical category/currency median."""
        try:
            if not re.fullmatch(r"\d{4}-\d{2}", month):
                raise ValueError
            year, month_num = (int(part) for part in month.split("-"))
            datetime(year, month_num, 1)
        except (ValueError, TypeError) as exc:
            raise ValueError("Month must use YYYY-MM") from exc
        if not 1 <= lookback_months <= 36 or minimum_samples < 1 or factor <= 1:
            raise ValueError("Lookback must be 1..36 months, samples positive, factor greater than 1")
        target_index = year * 12 + month_num - 1
        start_index = target_index - lookback_months
        start_year, start_month0 = divmod(start_index, 12)
        start = f"{start_year:04d}-{start_month0 + 1:02d}-01"
        end = f"{year:04d}-{month_num:02d}-01"
        end_index = target_index + 1
        end_year, end_month0 = divmod(end_index, 12)
        next_month = f"{end_year:04d}-{end_month0 + 1:02d}-01"
        with self._db() as db:
            baseline_rows = list(db.execute(
                "SELECT category,currency,amount FROM transactions WHERE user_id=? AND booked_on>=? AND booked_on<?",
                (user_id, start, end),
            ))
            current_rows = list(db.execute(
                "SELECT id,booked_on,amount,currency,merchant,category FROM transactions WHERE user_id=? AND booked_on>=? AND booked_on<? ORDER BY booked_on,id",
                (user_id, end, next_month),
            ))
        baselines: dict[tuple[str, str], list[Decimal]] = {}
        for row in baseline_rows:
            baselines.setdefault((row["category"], row["currency"]), []).append(Decimal(row["amount"]))
        candidates = []
        for row in current_rows:
            values = sorted(baselines.get((row["category"], row["currency"]), []))
            if len(values) < minimum_samples:
                continue
            middle = len(values) // 2
            median = values[middle] if len(values) % 2 else (values[middle - 1] + values[middle]) / 2
            amount = Decimal(row["amount"])
            if amount >= median * factor:
                candidates.append({
                    "transaction_id": row["id"], "date": row["booked_on"],
                    "merchant": row["merchant"], "category": row["category"],
                    "amount": amount, "currency": row["currency"],
                    "baseline_median": median, "baseline_count": len(values),
                    "factor": factor,
                })
        return candidates

    def export_csv(self,user_id:int)->str:
        with self._db() as db: rows=list(db.execute("SELECT booked_on,amount,currency,merchant,category,description,source,external_id FROM transactions WHERE user_id=? ORDER BY booked_on,id",(user_id,)))
        out=io.StringIO(); writer=csv.writer(out); writer.writerow(["date","amount","currency","merchant","category","description","source","external_id"])
        writer.writerows([tuple(row) for row in rows]); return out.getvalue()

    def backup(self, destination: str | Path) -> Path:
        """Create a consistent SQLite backup without overwriting an existing file."""
        target = Path(destination).expanduser().resolve()
        source = Path(self.path).expanduser().resolve()
        if target == source:
            raise ValueError("Backup destination must differ from the active database")
        if target.exists():
            raise FileExistsError(f"Backup destination already exists: {target}")
        target.parent.mkdir(parents=True, exist_ok=True)
        backup_db = sqlite3.connect(target)
        try:
            with self._db() as source_db:
                source_db.backup(backup_db)
        except Exception:
            backup_db.close()
            target.unlink(missing_ok=True)
            raise
        else:
            backup_db.close()
            target.chmod(0o600)
        return target

    def delete_user(self,user_id:int)->None:
        with self._db() as db:
            db.execute("DELETE FROM recurring_executions WHERE rule_id IN (SELECT id FROM recurring_rules WHERE user_id=?)",(user_id,))
            db.execute("DELETE FROM transactions WHERE user_id=?",(user_id,)); db.execute("DELETE FROM categories WHERE user_id=?",(user_id,)); db.execute("DELETE FROM monthly_budgets WHERE user_id=?",(user_id,)); db.execute("DELETE FROM recurring_rules WHERE user_id=?",(user_id,)); db.execute("DELETE FROM import_batches WHERE user_id=?",(user_id,))


@dataclass
class ImportPreview:
    rows:list[TransactionInput]
    errors:list[str]
    duplicates:int
    file_hash:str


def parse_statement(data:bytes, mapping:dict[str,str], filename:str)->ImportPreview:
    digest=hashlib.sha256(data).hexdigest(); parsed=[]; errors=[]
    try:
        if filename.lower().endswith((".xlsx",".xlsm")):
            from openpyxl import load_workbook
            sheet=load_workbook(io.BytesIO(data),read_only=True,data_only=True).active
            all_rows=sheet.iter_rows(values_only=True); headers=[str(x or "").strip() for x in next(all_rows)]
            records=[dict(zip(headers,r)) for r in all_rows]
        else:
            text=data.decode("utf-8-sig")
            sample=text[:4096]; dialect=csv.Sniffer().sniff(sample,delimiters=",;\t")
            records=list(csv.DictReader(io.StringIO(text),dialect=dialect))
    except (OSError, ValueError, TypeError, UnicodeError, csv.Error, StopIteration, zipfile.BadZipFile, InvalidFileException) as exc:
        return ImportPreview([], [f"Cannot read file: {type(exc).__name__}"],0,digest)
    required={"date","amount","merchant"}
    if not required.issubset(mapping):
        return ImportPreview([], ["Mapping must include date, amount and merchant"],0,digest)
    for line,record in enumerate(records,start=2):
        try:
            values={key:record.get(column) for key,column in mapping.items()}
            day=date.fromisoformat(str(values["date"]).strip()[:10])
            curr=str(values.get("currency") or "RUB").upper().strip()
            amt=money(str(values["amount"]).replace(" ","").replace(",","."),curr)
            merchant=str(values["merchant"] or "").strip()
            if not merchant: raise ValueError("merchant is empty")
            parsed.append(TransactionInput(day,amt,curr,merchant,str(values.get("category") or "Uncategorized").strip(),str(values.get("description") or "").strip(),"import",str(values.get("external_id") or "").strip() or None))
        except (ValueError,InvalidOperation,TypeError) as exc: errors.append(f"Row {line}: {exc}")
    seen=set(); unique=[]; duplicates=0
    for row in parsed:
        key=Ledger.fingerprint(row)
        if key in seen: duplicates+=1
        else: seen.add(key); unique.append(row)
    return ImportPreview(unique,errors,duplicates,digest)


def commit_import(ledger:Ledger,user_id:int,preview:ImportPreview)->dict:
    """Commit only after explicit caller confirmation; batch record makes reruns visible/idempotent."""
    imported=duplicates=0
    with ledger._db() as db:
        try: cur=db.execute("INSERT INTO import_batches(user_id,file_hash) VALUES(?,?)",(user_id,preview.file_hash))
        except sqlite3.IntegrityError: return {"batch_id":None,"imported":0,"duplicates":len(preview.rows),"errors":preview.errors,"already_imported":True}
        batch_id=int(cur.lastrowid)
        for row in preview.rows:
            if not row.merchant.strip(): raise ValueError("Merchant cannot be empty")
            try:
                db.execute("INSERT INTO transactions(user_id,booked_on,amount,currency,merchant,category,description,source,external_id,fingerprint,import_batch_id) VALUES(?,?,?,?,?,?,?,?,?,?,?)",
                    (user_id,row.booked_on.isoformat(),str(money(row.amount,row.currency)),row.currency.upper(),row.merchant.strip(),row.category.strip() or "Uncategorized",row.description,row.source,row.external_id,ledger.fingerprint(row),batch_id))
                imported+=1
            except sqlite3.IntegrityError: duplicates+=1
    return {"batch_id":batch_id,"imported":imported,"duplicates":duplicates+preview.duplicates,"errors":preview.errors,"already_imported":False}


@dataclass(frozen=True)
class ExchangeRateQuote:
    currency: str
    requested_date: date
    effective_date: date
    rate: Decimal


class CbrRates:
    """Public Bank of Russia XML adapter with date-keyed cache and effective-date metadata."""
    BASE="https://www.cbr.ru/scripts/"
    def __init__(self,transport:Callable[[str,float],bytes]|None=None,timeout:float=5.0,ledger:Ledger|None=None):
        self.transport=transport or self._get; self.timeout=timeout; self.ledger=ledger
        self.cache:dict[tuple[str,date],ExchangeRateQuote]={}
    @staticmethod
    def _get(url:str,timeout:float)->bytes:
        request=urllib.request.Request(url,headers={"User-Agent":"MoneyRadar/0.1 (+personal local tool)"})
        context=ssl.create_default_context(cafile=certifi.where())
        with urllib.request.urlopen(request,timeout=timeout,context=context) as response: return response.read(2_000_000)
    @staticmethod
    def _effective_date(value:str|None,default:date)->date:
        if not value:return default
        try:return datetime.strptime(value,"%d.%m.%Y").date()
        except ValueError as exc:raise RuntimeError("CBR returned an invalid effective date") from exc

    def quote(self,currency:str,on_date:date|None=None)->ExchangeRateQuote:
        code=currency.upper(); day=on_date or date.today(); key=(code,day)
        if not re.fullmatch(r"[A-Z]{3}", code):
            raise ValueError("Currency must be a three-letter ISO code")
        if key in self.cache:return self.cache[key]
        if self.ledger:
            cached=self.ledger.cached_exchange_rate(code,day)
            if cached:
                quote=ExchangeRateQuote(code,day,cached[1],cached[0]);self.cache[key]=quote;return quote
        if code=="RUB":
            quote=ExchangeRateQuote(code,day,day,Decimal(1))
            if self.ledger:self.ledger.save_exchange_rate(code,day,day,quote.rate)
            self.cache[key]=quote;return quote
        url=f"{self.BASE}XML_daily.asp?date_req={day:%d/%m/%Y}"
        try: payload=self.transport(url,self.timeout)
        except (TimeoutError,urllib.error.URLError) as exc: raise RuntimeError("CBR rates are temporarily unavailable; retry later") from exc
        root=ET.fromstring(payload)
        if on_date:
            ref=next((n.attrib.get("ID") for n in root.findall("Valute") if n.findtext("CharCode")==code),None)
            if ref is None: raise ValueError(f"Currency {code} not found in CBR rates for {day}")
            url=f"{self.BASE}XML_dynamic.asp?date_req1={day:%d/%m/%Y}&date_req2={day:%d/%m/%Y}&VAL_NM_RQ={ref}"
            try: payload=self.transport(url,self.timeout)
            except (TimeoutError,urllib.error.URLError) as exc: raise RuntimeError("CBR rates are temporarily unavailable; retry later") from exc
            root=ET.fromstring(payload)
            vals=root.findall("Record")
            if not vals: raise ValueError(f"No CBR rate available for {code} on {day}")
            node=vals[-1]
            effective=self._effective_date(node.attrib.get("Date"),day)
            if effective!=day:raise ValueError(f"No exact CBR rate available for {code} on {day}")
            nominal=Decimal(node.findtext("Nominal","1")); value=Decimal(node.findtext("Value","0").replace(",",".")); result=value/nominal
        else:
            node=next((n for n in root.findall("Valute") if n.findtext("CharCode")==code),None)
            if node is None: raise ValueError(f"Currency {code} not found in CBR daily rates")
            effective=self._effective_date(root.attrib.get("Date"),day)
            result=Decimal(node.findtext("Value","0").replace(",","."))/Decimal(node.findtext("Nominal","1"))
        quote=ExchangeRateQuote(code,day,effective,result)
        if self.ledger:self.ledger.save_exchange_rate(code,day,effective,result)
        self.cache[key]=quote;return quote

    def rate(self,currency:str,on_date:date|None=None)->Decimal:
        return self.quote(currency,on_date).rate
